from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = "docs/eventflow-data-protection-report.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9E2EF")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(46, 116, 181 if level < 3 else 120)
    return paragraph


def add_body(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    return paragraph


def add_status_table(document, rows):
    table = document.add_table(rows=1, cols=4)
    set_table_width(table, [1.5, 2.0, 1.2, 1.8])
    set_table_borders(table)
    headers = ["Area", "Implementation", "Status", "Evidence"]
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[idx], "2E74B5")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
        if row[2] == "Done":
            set_cell_shading(cells[2], "E7F4EA")
        elif row[2] == "Partial":
            set_cell_shading(cells[2], "FFF4D6")
        else:
            set_cell_shading(cells[2], "FDEAEA")
    document.add_paragraph()


def add_commit_plan_table(document):
    rows = [
        [
            "feat(security): harden auth session and consent controls",
            "Backend and frontend auth changes, consent storage, public legal pages.",
            "Keeps auth changes reviewable without mixing report artifacts.",
        ],
        [
            "feat(data-protection): add user data export, erasure, retention, and payment minimization",
            "User control APIs, anonymization path, retention scheduler, payment payload sanitizer, migration.",
            "Separates privacy rights and retention behavior from session mechanics.",
        ],
        [
            "docs(data-protection): add lecturer report and verification notes",
            "DOCX report plus generator used to recreate the artifact.",
            "Makes the submitted artifact traceable and maintainable.",
        ],
    ]
    table = document.add_table(rows=1, cols=3)
    set_table_width(table, [2.4, 2.3, 1.8])
    set_table_borders(table)
    headers = ["Commit", "Scope", "Reason"]
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[idx], "1F4D78")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    document.add_paragraph()


def build_document():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("EventFlow Data Protection Implementation Report")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(31, 77, 120)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Prepared for lecturer review - July 03, 2026")
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(85, 85, 85)

    add_body(
        document,
        "This report summarizes the Data Protection controls added to EventFlow MVP. "
        "The goal is to show what the project now guarantees technically, how user rights are supported, "
        "and what remains as future improvement before production-grade legal compliance.",
    )

    add_heading(document, "1. Executive Summary", 1)
    add_body(
        document,
        "EventFlow now has a stronger baseline for confidentiality, integrity, availability, privacy, "
        "and user control. The implementation focuses on enforceable consent, safer session storage, "
        "data export and erasure rights, retention cleanup, payment data minimization, production endpoint hardening, "
        "and a clearer privacy/terms surface.",
    )
    add_bullet(document, "Authentication tokens are issued through httpOnly cookies instead of exposing new tokens to localStorage.")
    add_bullet(document, "Signup now requires explicit acceptance of Privacy Policy and Terms of Service.")
    add_bullet(document, "Users can export their personal data and request anonymization from the Profile page.")
    add_bullet(document, "Security and retention controls are backed by automated tests, linting, and production build checks.")

    add_heading(document, "2. Data Protection Scope", 1)
    add_status_table(
        document,
        [
            ["Confidentiality", "httpOnly auth cookies, sanitized payment payloads, no password/token response body after login.", "Done", "AuthController, JwtAuthFilter, PayOsPaymentService"],
            ["Integrity", "Server-side consent validation, JWT validation, account-erasure guard for deleted users.", "Done", "SignupRequest, AuthService, JwtAuthFilter"],
            ["Availability", "Backup runbook retained; retention cleanup avoids unbounded security-data growth.", "Partial", "deployment runbook, DataRetentionScheduler"],
            ["Privacy", "Privacy/Terms pages, explicit consent capture, payment data minimization.", "Done", "LegalPage, User consent fields"],
            ["Control", "User data export and personal-data erasure endpoints with UI controls.", "Done", "UserController, UserProfileService, ProfilePage"],
        ],
    )

    add_heading(document, "3. Implemented Controls", 1)
    add_heading(document, "3.1 Legal and Consent Surface", 2)
    add_bullet(document, "Created public Privacy Policy and Terms pages at /privacy and /terms.")
    add_bullet(document, "Added signup checkbox linking directly to both legal pages.")
    add_bullet(document, "Stored consent version and consent timestamp on the user record.")
    add_bullet(document, "Rejected signup requests server-side when consent is missing or false.")

    add_heading(document, "3.2 Authentication and Session Protection", 2)
    add_bullet(document, "Login, verify-email, and refresh responses now set access and refresh tokens as httpOnly cookies.")
    add_bullet(document, "New auth responses no longer expose token values in the JSON response body.")
    add_bullet(document, "Frontend uses withCredentials and removes new token persistence in localStorage.")
    add_bullet(document, "JWT filter accepts cookie tokens and rejects users whose personal data has been erased.")

    add_heading(document, "3.3 User Rights: Export and Erasure", 2)
    add_bullet(document, "Added GET /users/{userId}/data-export for profile, consent, event membership, and payment summary.")
    add_bullet(document, "Added DELETE /users/{userId}/personal-data to anonymize account identity and revoke refresh tokens.")
    add_bullet(document, "Profile page now includes data export and personal-data deletion actions with confirmation.")
    add_bullet(document, "Erasure preserves business history without retaining direct account identifiers.")

    add_heading(document, "3.4 Retention and Data Minimization", 2)
    add_bullet(document, "Added scheduled cleanup for expired refresh tokens, old revoked refresh tokens, and old audit logs.")
    add_bullet(document, "Added configurable retention values through application.yml and production env example.")
    add_bullet(document, "Sanitized payOS payload persistence to avoid storing raw provider response bodies.")
    add_bullet(document, "Production env example now sets secure auth cookies and blocks public operations endpoints.")

    add_heading(document, "4. Verification Evidence", 1)
    add_status_table(
        document,
        [
            ["Backend tests", "mvn test", "Done", "33 tests, 0 failures"],
            ["Frontend lint", "npm run lint", "Done", "ESLint passed"],
            ["Frontend build", "npm run build", "Done", "Vite production build passed"],
            ["Diff hygiene", "git diff --check", "Done", "No whitespace/conflict-marker errors"],
            ["GitNexus review", "detect_changes", "Done", "High risk due to auth/user/security scope; reviewed intentionally"],
        ],
    )

    add_heading(document, "5. Maintenance and Scaling Notes", 1)
    add_bullet(document, "Keep auth/session changes isolated from UI-only commits to simplify future regression review.")
    add_bullet(document, "Version future policy updates by changing consentVersion and requiring renewed consent if policy scope changes materially.")
    add_bullet(document, "Move field-level encryption for phone, Telegram ID, attendee email/phone into a separate migration when key management is ready.")
    add_bullet(document, "Add admin-facing DPA/vendor inventory before connecting more third-party providers.")
    add_bullet(document, "Test restore drills and document incident-response steps before production launch.")

    add_heading(document, "6. Remaining Gaps", 1)
    add_status_table(
        document,
        [
            ["Legal review", "Privacy/Terms text should be reviewed by a qualified legal owner.", "Partial", "Current text is engineering-ready, not legal advice"],
            ["Field encryption", "PII fields are minimized and controllable but not yet encrypted field-by-field.", "Partial", "Future encryption service needed"],
            ["Incident response", "Runbook should define breach triage, notification owner, and timeline.", "Partial", "Process document needed"],
            ["Vendor governance", "DPA/vendor register should be formalized for AI, email, Telegram, storage, and payment.", "Partial", "Operational checklist needed"],
        ],
    )

    add_heading(document, "7. Commit Plan", 1)
    add_commit_plan_table(document)

    add_heading(document, "Conclusion", 1)
    add_body(
        document,
        "The project is now substantially stronger than the initial MVP state for Data Protection. "
        "It includes concrete user-facing rights, safer authentication storage, clearer consent, retention cleanup, "
        "and reduced payment-data exposure. The remaining work is mainly production governance: legal review, "
        "field-level encryption, DPA management, and incident-response procedure.",
    )

    document.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
